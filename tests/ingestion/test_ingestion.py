import csv
import json
import tempfile
import unittest
from pathlib import Path
from unittest import mock

import validar_documentos
from openpyxl import Workbook

import lector_corpus as lector
from nerv.ingestion import lector_corpus as packaged_lector


class FakePage:
    def __init__(self, text):
        self.text = text

    def get_text(self, _mode):
        return self.text


class FakeDocument:
    def __init__(self, pages, needs_pass=False):
        self.pages = pages
        self.needs_pass = needs_pass
        self.page_count = len(pages)

    def __enter__(self):
        return self

    def __exit__(self, *_args):
        return False

    def __iter__(self):
        return iter(self.pages)


class FakeFitz:
    def __init__(self, document):
        self.document = document

    def open(self, _archivo):
        return self.document


def enc_varint(value):
    salida = bytearray()

    while True:
        byte = value & 0x7F
        value >>= 7

        if value:
            salida.append(byte | 0x80)
        else:
            salida.append(byte)
            return bytes(salida)


def campo_varint(numero, value):
    return enc_varint((numero << 3) | 0) + enc_varint(value)


def campo_bytes(numero, value):
    return enc_varint((numero << 3) | 2) + enc_varint(len(value)) + value


def campo_string(numero, value):
    return campo_bytes(numero, value.encode("utf-8"))


def tile_mvt_minima():
    valor_nombre = campo_string(1, "Puerto Esperanza")
    valor_riesgo = campo_string(1, "alto")
    tags = enc_varint(0) + enc_varint(0) + enc_varint(1) + enc_varint(1)
    feature = (
        campo_varint(1, 7)
        + campo_bytes(2, tags)
        + campo_varint(3, 1)
    )
    layer = (
        campo_string(1, "observatorio")
        + campo_bytes(2, feature)
        + campo_string(3, "nombre")
        + campo_string(3, "riesgo")
        + campo_bytes(4, valor_nombre)
        + campo_bytes(4, valor_riesgo)
        + campo_varint(5, 4096)
        + campo_varint(15, 2)
    )

    return campo_bytes(3, layer)


class IngestaTests(unittest.TestCase):
    def test_cli_compatibility_wrapper_uses_packaged_implementation(self):
        self.assertIs(lector, packaged_lector)

    def test_limpiar_texto_conserva_parrafos_y_reduce_ruido(self):
        texto = " Uno\t dos \r\n\r\n\r\n tres\x00 "

        self.assertEqual(
            lector.limpiar_texto(texto),
            "Uno dos\n\ntres"
        )

    def test_obtener_formato_normaliza_extensiones(self):
        self.assertEqual(lector.obtener_formato(".htm"), "html")
        self.assertEqual(lector.obtener_formato(".jpg"), "imagen")
        self.assertEqual(lector.obtener_formato(".pbf"), "pbf")
        self.assertEqual(lector.obtener_formato(".csv"), "csv")

    def test_obtener_fenomeno_desde_carpetas(self):
        with tempfile.TemporaryDirectory() as tmp:
            corpus = Path(tmp)
            archivo = corpus / "F3_demo" / "a.txt"
            archivo.parent.mkdir()
            archivo.write_text("x", encoding="utf-8")

            self.assertEqual(lector.obtener_fenomeno(archivo, corpus), 3)

    def test_doc_id_estable_reproducible_e_incluye_ruta(self):
        with tempfile.TemporaryDirectory() as tmp:
            corpus = Path(tmp)
            a = corpus / "F1_demo" / "a.txt"
            b = corpus / "F1_demo" / "b.txt"
            a.parent.mkdir()
            a.write_text("mismo contenido", encoding="utf-8")
            b.write_text("mismo contenido", encoding="utf-8")

            self.assertEqual(
                lector.generar_doc_id(a, corpus),
                lector.generar_doc_id(a, corpus)
            )
            self.assertNotEqual(
                lector.generar_doc_id(a, corpus),
                lector.generar_doc_id(b, corpus)
            )

    def test_fingerprint_rapido_no_depende_de_leer_contenido(self):
        with tempfile.TemporaryDirectory() as tmp:
            corpus = Path(tmp)
            archivo = corpus / "F1_demo" / "a.txt"
            archivo.parent.mkdir()
            archivo.write_text("contenido", encoding="utf-8")

            fingerprint = lector.generar_fingerprint_archivo(
                archivo,
                corpus,
                modo=lector.MODO_DOC_ID_RAPIDO
            )

        self.assertTrue(fingerprint["doc_id"].startswith("DOC-"))
        self.assertEqual(len(fingerprint["sha256"]), 64)

    def test_lectores_txt_json_csv_xlsx_html(self):
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            txt = base / "a.txt"
            txt.write_text("texto plano", encoding="utf-8")

            js = base / "a.json"
            js.write_text(
                json.dumps({"titulo": "Uno", "items": [{"valor": 2}]}),
                encoding="utf-8"
            )

            csv_path = base / "a.csv"
            with csv_path.open("w", encoding="utf-8", newline="") as fh:
                writer = csv.DictWriter(fh, fieldnames=["nombre", "valor"])
                writer.writeheader()
                writer.writerow({"nombre": "alfa", "valor": "10"})

            xlsx = base / "a.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "Datos"
            ws.append(["nombre", "valor"])
            ws.append(["beta", 20])
            wb.save(xlsx)

            html = base / "a.html"
            html.write_text(
                "<html><body><nav>menu</nav><main><h1>Titulo</h1>"
                "<p>Texto visible</p></main><script>x</script></body></html>",
                encoding="utf-8"
            )

            self.assertEqual(lector.leer_txt(txt), "texto plano")
            self.assertIn("titulo: Uno", lector.leer_json(js))
            self.assertIn("nombre: alfa | valor: 10", lector.leer_csv(csv_path))
            self.assertIn(
                "[Hoja: Datos] nombre: beta | valor: 20",
                lector.leer_xlsx(xlsx),
            )
            self.assertIn("Texto visible", lector.leer_html(html))
            self.assertNotIn("menu", lector.leer_html(html))

    def test_json_vacio_devuelve_texto_vacio(self):
        with tempfile.TemporaryDirectory() as tmp:
            ruta = Path(tmp) / "vacio.json"
            ruta.write_text("[]", encoding="utf-8")

            self.assertEqual(lector.leer_json(ruta), "")

    def test_json_utf8_con_bom_se_lee_correctamente(self):
        with tempfile.TemporaryDirectory() as tmp:
            ruta = Path(tmp) / "bom.json"
            ruta.write_text(
                json.dumps({"titulo": "Con BOM"}),
                encoding="utf-8-sig"
            )

            self.assertIn("titulo: Con BOM", lector.leer_json(ruta))

    def test_json_conserva_urls_como_metadata(self):
        with tempfile.TemporaryDirectory() as tmp:
            corpus = Path(tmp)
            carpeta = corpus / "F1_demo"
            carpeta.mkdir()
            ruta = carpeta / "fuente.json"
            ruta.write_text(
                json.dumps({
                    "titulo": "Informe",
                    "source": "https://example.org/informe.pdf",
                    "items": [
                        {"url": "https://datos.example.org/base?id=1"}
                    ]
                }),
                encoding="utf-8"
            )

            documentos, errores, estadisticas = lector.procesar_corpus(
                corpus,
                ruta_cache=None
            )

        self.assertEqual(errores, [])
        self.assertEqual(estadisticas["procesados"], 1)
        self.assertEqual(
            documentos[0]["metadata"]["procedencia"]["urls"],
            [
                "https://datos.example.org/base?id=1",
                "https://example.org/informe.pdf"
            ]
        )
        self.assertIn("titulo: Informe", documentos[0]["texto"])
        self.assertNotIn("https://example.org/informe.pdf", documentos[0]["texto"])

    def test_pdf_normal_no_usa_ocr(self):
        doc = FakeDocument([FakePage("Texto normal")])
        diagnostico = {}

        with mock.patch.object(lector, "fitz", FakeFitz(doc)):
            with mock.patch.object(
                lector,
                "_leer_ocr_pagina_pdf",
                side_effect=AssertionError("no debe usar OCR")
            ):
                texto = lector.leer_pdf(
                    Path("normal.pdf"),
                    diagnostico=diagnostico
                )

        self.assertIn("Texto normal", texto)
        self.assertEqual(diagnostico, {
            "paginas": 1,
            "paginas_nativas": 1,
            "paginas_ocr": 0,
            "paginas_sin_texto": 0
        })

    def test_texto_nativo_util_acepta_paginas_cortas(self):
        self.assertFalse(lector._texto_nativo_util(""))
        self.assertFalse(lector._texto_nativo_util("  \n "))
        self.assertTrue(lector._texto_nativo_util("A"))
        self.assertTrue(lector._texto_nativo_util("7"))

        doc = FakeDocument([FakePage("A")])

        with mock.patch.object(lector, "fitz", FakeFitz(doc)):
            with mock.patch.object(
                lector,
                "_leer_ocr_pagina_pdf",
                side_effect=AssertionError("no debe usar OCR")
            ):
                texto = lector.leer_pdf(Path("pagina-corta.pdf"))

        self.assertEqual(texto, "[Pagina 1]\nA")

    def test_pdf_escaneado_activa_ocr(self):
        doc = FakeDocument([FakePage("")])
        diagnostico = {}
        fitz_mock = mock.Mock()
        fitz_mock.open.return_value = doc

        with mock.patch.object(lector, "fitz", fitz_mock):
            with mock.patch.object(
                lector,
                "_leer_ocr_pagina_pdf",
                return_value="[Pagina 1]\nTexto OCR"
            ) as ocr:
                texto = lector.leer_pdf(
                    Path("scan.pdf"),
                    diagnostico=diagnostico
                )

        fitz_mock.open.assert_called_once_with(Path("scan.pdf"))
        ocr.assert_called_once()
        self.assertIn("Texto OCR", texto)
        self.assertEqual(diagnostico["paginas_nativas"], 0)
        self.assertEqual(diagnostico["paginas_ocr"], 1)
        self.assertEqual(diagnostico["paginas_sin_texto"], 0)

    def test_pdf_escaneado_sin_ocr_falla_rapido(self):
        doc = FakeDocument([FakePage("")])

        with mock.patch.object(lector, "fitz", FakeFitz(doc)):
            with mock.patch.object(
                lector,
                "_leer_ocr_pagina_pdf",
                side_effect=AssertionError("no debe usar OCR")
            ):
                with self.assertRaises(ValueError):
                    lector.leer_pdf(Path("scan.pdf"), usar_ocr=False)

    def test_pdf_mixto_aplica_ocr_solo_a_paginas_vacias(self):
        doc = FakeDocument([
            FakePage("Primera"),
            FakePage(""),
            FakePage("Tercera")
        ])
        diagnostico = {}

        with mock.patch.object(lector, "fitz", FakeFitz(doc)):
            with mock.patch.object(
                lector,
                "_leer_ocr_pagina_pdf",
                return_value="[Pagina 2]\nSegunda"
            ) as ocr:
                texto = lector.leer_pdf(
                    Path("mixto.pdf"),
                    diagnostico=diagnostico
                )

        ocr.assert_called_once()
        self.assertLess(texto.index("Primera"), texto.index("Segunda"))
        self.assertLess(texto.index("Segunda"), texto.index("Tercera"))
        self.assertEqual(texto.count("[Pagina 1]"), 1)
        self.assertEqual(texto.count("[Pagina 2]"), 1)
        self.assertEqual(texto.count("[Pagina 3]"), 1)
        self.assertEqual(diagnostico["paginas_nativas"], 2)
        self.assertEqual(diagnostico["paginas_ocr"], 1)
        self.assertEqual(diagnostico["paginas_sin_texto"], 0)

    def test_pdf_mixto_sin_ocr_se_difiere_completo(self):
        doc = FakeDocument([
            FakePage("Primera"),
            FakePage(""),
            FakePage("Tercera")
        ])

        with mock.patch.object(lector, "fitz", FakeFitz(doc)):
            with mock.patch.object(
                lector,
                "_leer_ocr_pagina_pdf",
                side_effect=AssertionError("no debe usar OCR")
            ):
                with self.assertRaisesRegex(
                    ValueError,
                    "PDF con paginas sin texto extraible"
                ):
                    lector.leer_pdf(
                        Path("mixto.pdf"),
                        usar_ocr=False
                    )

    def test_pdf_mixto_con_pagina_sin_texto_conserva_contenido_util(self):
        doc = FakeDocument([
            FakePage("Primera"),
            FakePage(""),
            FakePage("")
        ])
        diagnostico = {}

        with mock.patch.object(lector, "fitz", FakeFitz(doc)):
            with mock.patch.object(
                lector,
                "_leer_ocr_pagina_pdf",
                side_effect=["", "[Pagina 3]\nTercera"]
            ) as ocr:
                texto = lector.leer_pdf(
                    Path("mixto.pdf"),
                    diagnostico=diagnostico
                )

        self.assertEqual(ocr.call_count, 2)
        self.assertIn("Primera", texto)
        self.assertIn("Tercera", texto)
        self.assertNotIn("[Pagina 2]", texto)
        self.assertEqual(diagnostico["paginas_nativas"], 1)
        self.assertEqual(diagnostico["paginas_ocr"], 1)
        self.assertEqual(diagnostico["paginas_sin_texto"], 1)

    def test_procesar_corpus_publica_pdf_con_texto_util_parcial(self):
        with tempfile.TemporaryDirectory() as tmp:
            corpus = Path(tmp)
            carpeta = corpus / "F1_demo"
            carpeta.mkdir()
            pdf = carpeta / "mixto.pdf"
            pdf.write_bytes(b"%PDF")
            doc = FakeDocument([FakePage("Texto nativo"), FakePage("")])

            with mock.patch.object(lector, "fitz", FakeFitz(doc)):
                with mock.patch.object(
                    lector,
                    "_leer_ocr_pagina_pdf",
                    return_value=""
                ):
                    documentos, errores, estadisticas = lector.procesar_corpus(
                        corpus,
                        ruta_cache=None,
                        usar_ocr=True,
                        max_trabajadores=1,
                        max_procesos_pesados=1
                    )

        self.assertEqual(len(documentos), 1)
        self.assertIn("Texto nativo", documentos[0]["texto"])
        self.assertEqual(errores, [])
        self.assertEqual(estadisticas["procesados"], 1)
        self.assertEqual(estadisticas["fallidos"], 0)

    def test_pdf_mixto_con_excepcion_ocr_falla_claramente(self):
        doc = FakeDocument([FakePage("Texto nativo"), FakePage("")])

        with mock.patch.object(lector, "fitz", FakeFitz(doc)):
            with mock.patch.object(
                lector,
                "_leer_ocr_pagina_pdf",
                side_effect=ValueError("Fallo de OCR: runtime roto")
            ):
                with self.assertRaisesRegex(
                    ValueError,
                    "Fallo de OCR: runtime roto"
                ):
                    lector.leer_pdf(Path("mixto.pdf"))

    def test_imagen_con_texto_y_sin_texto_mediante_mock(self):
        with tempfile.TemporaryDirectory() as tmp:
            ruta = Path(tmp) / "imagen.png"
            lector.Image.new("RGB", (20, 20), "white").save(ruta)

            with mock.patch.object(lector, "_validar_ocr_disponible"):
                with mock.patch.object(
                    lector,
                    "_ocr_imagen",
                    return_value=" texto detectado "
                ):
                    self.assertEqual(
                        lector.leer_imagen(ruta),
                        "texto detectado"
                    )

                with mock.patch.object(
                    lector,
                    "_ocr_imagen",
                    return_value="   "
                ):
                    self.assertEqual(lector.leer_imagen(ruta), "")

    def test_ocr_normaliza_formato_de_imagen_para_tesseract(self):
        imagen = lector.Image.new("RGB", (20, 20), "white")
        imagen.format = "AVIF"

        with mock.patch.object(lector, "_validar_ocr_disponible"):
            with mock.patch.object(
                lector.pytesseract,
                "image_to_string",
                return_value="texto"
            ) as ocr:
                self.assertEqual(lector._ocr_imagen(imagen), "texto")

        imagen_enviada = ocr.call_args.args[0]
        self.assertIsNone(imagen_enviada.format)

    def test_imagen_sin_texto_conserva_metadata_en_incidencia(self):
        with tempfile.TemporaryDirectory() as tmp:
            corpus = Path(tmp)
            carpeta = corpus / "F3_demo"
            carpeta.mkdir()
            ruta = carpeta / "foto.jpg"
            lector.Image.new("RGB", (20, 10), "white").save(ruta)

            with mock.patch.object(lector, "_validar_ocr_disponible"):
                with mock.patch.object(
                    lector,
                    "_ocr_imagen",
                    return_value="   "
                ):
                    documentos, errores, estadisticas = (
                        lector.procesar_corpus(
                            corpus,
                            ruta_cache=None,
                            max_procesos_pesados=1
                        )
                    )

        self.assertEqual(documentos, [])
        self.assertEqual(errores[0]["estado"], "sin_texto_util")
        self.assertEqual(errores[0]["metadata"]["imagen"]["ancho"], 20)
        self.assertEqual(errores[0]["metadata"]["imagen"]["alto"], 10)
        self.assertEqual(estadisticas["vacios"], 1)

    def test_pbf_valido_duplicado_y_corrupto(self):
        with tempfile.TemporaryDirectory() as tmp:
            ruta = Path(tmp) / "tile.pbf"
            ruta.write_bytes(tile_mvt_minima())
            fingerprints = set()

            primero = lector.leer_pbf(ruta, fingerprints)
            segundo = lector.leer_pbf(ruta, fingerprints)
            diagnostico = lector.diagnosticar_pbf(ruta, fingerprints)

            corrupto = Path(tmp) / "corrupto.pbf"
            corrupto.write_bytes(b"\x80")

            with self.assertRaises(ValueError):
                lector.leer_pbf(corrupto, set())

        self.assertIn("capa: observatorio", primero)
        self.assertIn("nombre: Puerto Esperanza", primero)
        self.assertEqual(segundo, "")
        self.assertEqual(diagnostico["estado"], "VALID")
        self.assertEqual(diagnostico["features"], 1)
        self.assertEqual(diagnostico["features_duplicadas"], 1)

    def test_diagnosticar_pbf_resume_tile_valido_y_corrupto(self):
        with tempfile.TemporaryDirectory() as tmp:
            valido = Path(tmp) / "tile.pbf"
            valido.write_bytes(tile_mvt_minima())
            corrupto = Path(tmp) / "corrupto.pbf"
            corrupto.write_bytes(b"\x80")

            diagnostico_valido = lector.diagnosticar_pbf(valido)
            diagnostico_corrupto = lector.diagnosticar_pbf(corrupto)

        self.assertEqual(diagnostico_valido["estado"], "VALID")
        self.assertEqual(diagnostico_valido["capas"], 1)
        self.assertEqual(diagnostico_valido["features"], 1)
        self.assertEqual(diagnostico_corrupto["estado"], "PARSE_ERROR")

    def test_procesar_corpus_no_publica_pbf_duplicado(self):
        with tempfile.TemporaryDirectory() as tmp:
            corpus = Path(tmp)
            carpeta = corpus / "F3_demo"
            carpeta.mkdir()
            (carpeta / "a.pbf").write_bytes(tile_mvt_minima())
            (carpeta / "b.pbf").write_bytes(tile_mvt_minima())

            documentos, errores, estadisticas = lector.procesar_corpus(
                corpus,
                ruta_cache=None
            )

        self.assertEqual(len(documentos), 1)
        self.assertIn("capa: observatorio", documentos[0]["texto"])
        self.assertNotIn(
            "PBF valido sin contenido nuevo",
            documentos[0]["texto"]
        )
        self.assertEqual(len(errores), 1)
        self.assertEqual(errores[0]["archivo"], "F3_demo/b.pbf")
        self.assertEqual(errores[0]["estado"], "vacío")
        self.assertEqual(estadisticas["procesados"], 1)
        self.assertEqual(estadisticas["vacios"], 1)

    def test_procesar_corpus_contratos_basura_auxiliares_orden_y_fallos(self):
        with tempfile.TemporaryDirectory() as tmp:
            corpus = Path(tmp)
            (corpus / ".DS_Store").write_bytes(b"basura")
            (corpus / "Indice_Datos_Codefest.xlsx").write_bytes(b"")
            carpeta = corpus / "F1_demo"
            carpeta.mkdir()
            (carpeta / "b.txt").write_text(
                "el texto beta tiene palabras suficientes",
                encoding="utf-8"
            )
            (carpeta / "a.txt").write_text(
                "el texto alfa tiene palabras suficientes",
                encoding="utf-8"
            )
            (carpeta / "fallo.pbf").write_bytes(b"\x80")

            documentos, errores, estadisticas = lector.procesar_corpus(corpus)

        fuentes = [doc["fuente"] for doc in documentos]
        self.assertEqual(fuentes, sorted(fuentes))
        self.assertEqual(len(documentos), 2)
        self.assertEqual(estadisticas["procesados"], 2)
        self.assertEqual(estadisticas["omitidos"], 1)
        self.assertEqual(estadisticas["fallidos"], 1)

        for documento in documentos:
            self.assertFalse(Path(documento["fuente"]).is_absolute())
            self.assertFalse("\\" in documento["fuente"])
            self.assertFalse(
                validar_documentos.validar_documento(
                    documento,
                    1,
                    set()
                )
            )

        estados = {error["estado"] for error in errores}
        self.assertIn("omitido", estados)
        self.assertIn("fallido", estados)

    def test_procesar_corpus_excluye_evaluacion_anidada_sin_excluir_xlsx_normal(
        self,
    ):
        with tempfile.TemporaryDirectory() as tmp:
            corpus = Path(tmp)
            carpeta = corpus / "F3_Dinamicas_Territoriales"
            carpeta.mkdir()

            auxiliar = carpeta / "FASE ORDENADA CODEFEST.xlsx"
            wb_auxiliar = Workbook()
            wb_auxiliar.active.append(["PREGUNTA", "FRAGMENTO", "DOCUMENTO"])
            wb_auxiliar.active.append(["pregunta sintetica", "texto", "fuente"])
            wb_auxiliar.save(auxiliar)

            normal = carpeta / "fuente_legitima.xlsx"
            wb_normal = Workbook()
            wb_normal.active.title = "Datos"
            wb_normal.active.append(["nombre", "valor"])
            wb_normal.active.append(["fuente legitima", 1])
            wb_normal.save(normal)

            consulta_oficial = corpus / "Extracto_Preguntas_50_v2.pdf"
            consulta_oficial.write_bytes(b"contenido sintetico no parseable")

            documentos, errores, estadisticas = lector.procesar_corpus(
                corpus,
                ruta_cache=None,
            )

        self.assertEqual(
            [documento["fuente"] for documento in documentos],
            ["F3_Dinamicas_Territoriales/fuente_legitima.xlsx"],
        )
        self.assertIn("fuente legitima", documentos[0]["texto"])
        errores_por_archivo = {error["archivo"]: error for error in errores}
        self.assertEqual(
            errores_por_archivo[
                "F3_Dinamicas_Territoriales/FASE ORDENADA CODEFEST.xlsx"
            ]["estado"],
            "omitido",
        )
        self.assertIn(
            "auxiliar",
            errores_por_archivo[
                "F3_Dinamicas_Territoriales/FASE ORDENADA CODEFEST.xlsx"
            ]["error"],
        )
        self.assertEqual(
            errores_por_archivo["Extracto_Preguntas_50_v2.pdf"]["estado"],
            "omitido",
        )
        self.assertEqual(len(errores), 2)
        self.assertEqual(estadisticas["encontrados"], 3)
        self.assertEqual(estadisticas["procesados"], 1)
        self.assertEqual(estadisticas["omitidos"], 2)
        self.assertEqual(estadisticas["fallidos"], 0)

    def test_pdf_sin_ocr_se_registra_como_pendiente_ocr(self):
        with tempfile.TemporaryDirectory() as tmp:
            corpus = Path(tmp)
            carpeta = corpus / "F1_demo"
            carpeta.mkdir()
            pdf = carpeta / "scan.pdf"
            pdf.write_bytes(b"%PDF")
            doc = FakeDocument([FakePage("")])

            with mock.patch.object(lector, "fitz", FakeFitz(doc)):
                documentos, errores, estadisticas = lector.procesar_corpus(
                    corpus,
                    ruta_cache=None,
                    usar_ocr=False,
                    max_procesos_pesados=1
                )

        self.assertEqual(documentos, [])
        self.assertEqual(errores[0]["estado"], "pendiente_ocr")
        self.assertEqual(estadisticas["pendientes_ocr"], 1)
        self.assertEqual(estadisticas["fallidos"], 0)

    def test_pdf_mixto_sin_ocr_se_registra_como_pendiente_ocr(self):
        with tempfile.TemporaryDirectory() as tmp:
            corpus = Path(tmp)
            carpeta = corpus / "F1_demo"
            carpeta.mkdir()
            pdf = carpeta / "mixto.pdf"
            pdf.write_bytes(b"%PDF")
            doc = FakeDocument([
                FakePage("Texto nativo"),
                FakePage("")
            ])

            with mock.patch.object(lector, "fitz", FakeFitz(doc)):
                documentos, errores, estadisticas = lector.procesar_corpus(
                    corpus,
                    ruta_cache=None,
                    usar_ocr=False,
                    max_procesos_pesados=1
                )

        self.assertEqual(documentos, [])
        self.assertEqual(errores[0]["estado"], "pendiente_ocr")
        self.assertEqual(estadisticas["pendientes_ocr"], 1)
        self.assertEqual(estadisticas["procesados"], 0)

    def test_pdf_mixto_pendiente_se_resuelve_con_ocr_por_pagina(self):
        with tempfile.TemporaryDirectory() as tmp:
            corpus = Path(tmp)
            carpeta = corpus / "F1_demo"
            carpeta.mkdir()
            pdf = carpeta / "mixto.pdf"
            pdf.write_bytes(b"%PDF")
            doc = FakeDocument([
                FakePage("Texto nativo"),
                FakePage("")
            ])

            with mock.patch.object(lector, "fitz", FakeFitz(doc)):
                documentos, errores, estadisticas = lector.procesar_corpus(
                    corpus,
                    ruta_cache=None,
                    usar_ocr=False,
                    max_procesos_pesados=1
                )

                with mock.patch.object(
                    lector,
                    "_leer_ocr_pagina_pdf",
                    return_value="[Pagina 2]\nTexto OCR"
                ) as ocr:
                    documentos, errores, estadisticas = (
                        lector.resolver_pendientes_ocr(
                            corpus,
                            documentos,
                            errores,
                            {"F1_demo/mixto.pdf"},
                            estadisticas,
                            max_trabajadores=1,
                            max_procesos_pesados=1,
                            modo_doc_id=lector.MODO_DOC_ID_RAPIDO,
                            ruta_cache=None
                        )
                    )

        ocr.assert_called_once()
        self.assertEqual(errores, [])
        self.assertEqual(len(documentos), 1)
        self.assertLess(
            documentos[0]["texto"].index("Texto nativo"),
            documentos[0]["texto"].index("Texto OCR")
        )
        self.assertEqual(estadisticas["procesados"], 1)
        self.assertEqual(estadisticas["pendientes_ocr"], 0)

    def test_guardar_jsonl_salida_valida_sin_ids_duplicados(self):
        with tempfile.TemporaryDirectory() as tmp:
            corpus = Path(tmp) / "corpus"
            salida = Path(tmp) / "salida"
            carpeta = corpus / "F2_demo"
            carpeta.mkdir(parents=True)
            (carpeta / "a.txt").write_text(
                "the document has enough english words for detection",
                encoding="utf-8"
            )

            documentos, _errores, _estadisticas = lector.procesar_corpus(corpus)
            salida.mkdir()
            ruta = salida / "documentos.jsonl"
            lector.guardar_jsonl(documentos, ruta)

            ids = set()
            with ruta.open("r", encoding="utf-8") as fh:
                lineas = [json.loads(linea) for linea in fh]

        self.assertEqual(len(lineas), 1)
        errores = validar_documentos.validar_documento(lineas[0], 1, ids)
        self.assertEqual(errores, [])
        self.assertEqual(len(ids), 1)

    def test_procesar_corpus_reutiliza_cache_de_lecturas(self):
        with tempfile.TemporaryDirectory() as tmp:
            corpus = Path(tmp) / "corpus"
            carpeta = corpus / "F2_demo"
            carpeta.mkdir(parents=True)
            archivo = carpeta / "a.txt"
            archivo.write_text(
                "the document has enough english words for detection",
                encoding="utf-8"
            )
            ruta_cache = Path(tmp) / "cache" / "lecturas.jsonl"

            documentos, _errores, _estadisticas = lector.procesar_corpus(
                corpus,
                ruta_cache=ruta_cache
            )

            with mock.patch.object(
                lector,
                "leer_archivo",
                side_effect=AssertionError("debe usar cache")
            ):
                documentos_cache, _errores, _estadisticas = (
                    lector.procesar_corpus(
                        corpus,
                        ruta_cache=ruta_cache
                    )
                )

        self.assertEqual(documentos_cache, documentos)

    def test_corrida_rapida_reutiliza_cache_ocr_de_pdf(self):
        with tempfile.TemporaryDirectory() as tmp:
            corpus = Path(tmp) / "corpus"
            carpeta = corpus / "F1_demo"
            carpeta.mkdir(parents=True)
            archivo = carpeta / "scan.pdf"
            archivo.write_bytes(b"%PDF")
            ruta_cache = Path(tmp) / "cache" / "lecturas.jsonl"
            firma = lector.obtener_firma_cache(archivo, corpus)
            firma["usar_ocr"] = True
            doc_id = lector.generar_fingerprint_archivo(
                archivo,
                corpus
            )["doc_id"]
            cache = {
                lector.clave_cache(firma): {
                    "texto": "[Pagina 1]\nTexto OCR",
                    "doc_id": doc_id,
                    "doc_id_modo": lector.MODO_DOC_ID_RAPIDO
                }
            }
            lector.guardar_cache_lecturas(cache, ruta_cache)

            with mock.patch.object(
                lector,
                "leer_archivo",
                side_effect=AssertionError("debe usar cache OCR")
            ):
                documentos, errores, estadisticas = lector.procesar_corpus(
                    corpus,
                    ruta_cache=ruta_cache,
                    usar_ocr=False,
                    max_procesos_pesados=1
                )

        self.assertEqual(errores, [])
        self.assertEqual(len(documentos), 1)
        self.assertEqual(estadisticas["pendientes_ocr"], 0)

    def test_pendientes_ocr_ignoran_cache_rapido(self):
        with tempfile.TemporaryDirectory() as tmp:
            corpus = Path(tmp) / "corpus"
            carpeta = corpus / "F1_demo"
            carpeta.mkdir(parents=True)
            archivo = carpeta / "scan.pdf"
            archivo.write_bytes(b"%PDF")
            ruta_cache = Path(tmp) / "cache" / "lecturas.jsonl"
            firma = lector.obtener_firma_cache(archivo, corpus)
            firma["usar_ocr"] = False
            doc_id = lector.generar_fingerprint_archivo(
                archivo,
                corpus
            )["doc_id"]
            cache = {
                lector.clave_cache(firma): {
                    "texto": "[Pagina 1]\nTexto rapido",
                    "doc_id": doc_id,
                    "doc_id_modo": lector.MODO_DOC_ID_RAPIDO
                }
            }
            lector.guardar_cache_lecturas(cache, ruta_cache)

            with mock.patch.object(
                lector,
                "leer_archivo",
                return_value="[Pagina 1]\nTexto OCR"
            ) as leer:
                documentos, _errores, _estadisticas = lector.procesar_corpus(
                    corpus,
                    ruta_cache=ruta_cache,
                    usar_ocr=True,
                    ignorar_cache_lectura=True,
                    max_procesos_pesados=1
                )

        leer.assert_called_once()
        self.assertEqual(documentos[0]["texto"], "[Pagina 1]\nTexto OCR")

    def test_cargar_rutas_pendientes_ocr(self):
        with tempfile.TemporaryDirectory() as tmp:
            ruta = Path(tmp) / "errores.jsonl"
            lector.guardar_jsonl([
                {
                    "archivo": "F1/a.pdf",
                    "estado": "pendiente_ocr",
                    "error": "PDF sin texto extraible; OCR desactivado"
                },
                {
                    "archivo": "F1/b.pdf",
                    "estado": "fallido",
                    "error": "otro error"
                }
            ], ruta)

            pendientes = lector.cargar_rutas_pendientes_ocr(ruta)

        self.assertEqual(pendientes, {"F1/a.pdf"})

    def test_fusionar_por_fuente_reemplaza_documentos(self):
        base = [
            {"fuente": "b.txt", "texto": "viejo"},
            {"fuente": "a.txt", "texto": "base"}
        ]
        nuevos = [
            {"fuente": "b.txt", "texto": "nuevo"}
        ]

        fusionados = lector.fusionar_por_fuente(base, nuevos)

        self.assertEqual(
            [documento["fuente"] for documento in fusionados],
            ["a.txt", "b.txt"]
        )
        self.assertEqual(fusionados[1]["texto"], "nuevo")


if __name__ == "__main__":
    unittest.main()
