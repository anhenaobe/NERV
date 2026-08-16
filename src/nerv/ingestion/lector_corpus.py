import argparse
import csv
import hashlib
import json
import os
import re
import sys
import unicodedata
from concurrent.futures import ProcessPoolExecutor, ThreadPoolExecutor, as_completed
from pathlib import Path

try:
    from bs4 import BeautifulSoup
except ImportError:
    BeautifulSoup = None

try:
    import fitz
except ImportError:
    fitz = None

from openpyxl import load_workbook

try:
    from PIL import Image, UnidentifiedImageError
except ImportError:
    Image = None
    UnidentifiedImageError = OSError

try:
    import pytesseract
    from pytesseract import TesseractError, TesseractNotFoundError
except ImportError:
    pytesseract = None
    TesseractError = RuntimeError
    TesseractNotFoundError = RuntimeError

if pytesseract is not None:
    ruta_tesseract_windows = Path(
        r"C:\Program Files\Tesseract-OCR\tesseract.exe"
    )

    if ruta_tesseract_windows.exists():
        pytesseract.pytesseract.tesseract_cmd = str(
            ruta_tesseract_windows
        )


EXTENSIONES_IMAGEN = {
    ".jpg",
    ".jpeg",
    ".png",
    ".tif",
    ".tiff",
    ".webp",
    ".avif"
}

OCR_IDIOMAS = "spa+eng+por+chi_sim"
OCR_DPI = 220
RAIZ_PROYECTO = Path(__file__).resolve().parents[3]
TESSDATA_LOCAL = RAIZ_PROYECTO / "tessdata"
RUTA_CACHE_LECTURAS = RAIZ_PROYECTO / "corpus" / "cache" / "lecturas.jsonl"
CACHE_VERSION = 1
MODO_DOC_ID_RAPIDO = "rapido"
MODO_DOC_ID_CONTENIDO = "contenido"

BASURA_SILENCIOSA = {
    ".DS_Store",
    "Thumbs.db"
}

DIRECTORIOS_IGNORADOS = {
    "__pycache__",
    ".pytest_cache",
    ".venv",
    "venv",
    ".git"
}

ARCHIVOS_AUXILIARES = {
    "indice_datos_codefest.xlsx":
        "Archivo indice auxiliar del corpus",
    "extracto_preguntas_50_v2.pdf":
        "Archivo auxiliar de preguntas del corpus"
}

RUTAS_AUXILIARES = {
    "f3_dinamicas_territoriales/fase ordenada codefest.xlsx":
        "Archivo auxiliar de evaluacion del corpus"
}

EXTENSIONES_SOPORTADAS = {
    ".pdf",
    ".csv",
    ".json",
    ".html",
    ".htm",
    ".txt",
    ".md",
    ".xlsx",
    ".pbf",
    *EXTENSIONES_IMAGEN
}

MAX_TRABAJADORES = max(
    1,
    min(6, (os.cpu_count() or 1))
)

MAX_PROCESOS_PESADOS = max(
    1,
    min(2, (os.cpu_count() or 1))
)

EXTENSIONES_PROCESOS = {
    ".pdf",
    *EXTENSIONES_IMAGEN
}

PATRON_URL = re.compile(r"https?://[^\s<>'\")\]]+")


def _normalizar_url(url):
    return url.strip().rstrip(".,;)")


def _es_url_pura(valor):
    if not isinstance(valor, str):
        return False

    valor = valor.strip()

    if not valor:
        return False

    coincidencias = [
        _normalizar_url(url)
        for url in PATRON_URL.findall(valor)
    ]

    return len(coincidencias) == 1 and coincidencias[0] == valor


# =========================================================
# LIMPIEZA Y NORMALIZACIÓN
# =========================================================

def limpiar_texto(texto):
    texto = texto.replace("\r\n", "\n")
    texto = texto.replace("\r", "\n")
    texto = texto.replace("\t", " ")

    # Eliminar caracteres de control, conservando saltos de línea
    texto = "".join(
        caracter
        for caracter in texto
        if caracter == "\n" or caracter.isprintable()
    )

    # Reducir espacios repetidos
    texto = re.sub(r"[^\S\n]+", " ", texto)

    # Eliminar espacios alrededor de saltos de línea
    texto = re.sub(r" *\n *", "\n", texto)

    # Reducir tres o más saltos a dos
    texto = re.sub(r"\n{3,}", "\n\n", texto)

    return texto.strip()


def obtener_formato(extension):
    if extension in [".html", ".htm"]:
        return "html"

    if extension in EXTENSIONES_IMAGEN:
        return "imagen"

    if extension == ".pbf":
        return "pbf"

    return extension.lstrip(".").lower()


# =========================================================
# IDENTIFICACIÓN DOCUMENTAL
# =========================================================

def generar_doc_id(archivo, carpeta_corpus):
    return generar_fingerprint_archivo(
        archivo,
        carpeta_corpus,
        modo=MODO_DOC_ID_CONTENIDO
    )["doc_id"]


def generar_fingerprint_archivo(
    archivo,
    carpeta_corpus,
    modo=MODO_DOC_ID_RAPIDO
):
    hash_archivo = hashlib.sha256()

    ruta_relativa = archivo.relative_to(
        carpeta_corpus
    ).as_posix().lower()

    # Incluir la ruta evita que dos archivos idénticos
    # sean tratados como el mismo documento.
    hash_archivo.update(ruta_relativa.encode("utf-8"))

    if modo == MODO_DOC_ID_RAPIDO:
        estadistica = archivo.stat()
        hash_archivo.update(str(estadistica.st_size).encode("ascii"))
        hash_archivo.update(str(estadistica.st_mtime_ns).encode("ascii"))
    else:
        with open(archivo, "rb") as archivo_binario:
            for bloque in iter(
                lambda: archivo_binario.read(1024 * 1024),
                b""
            ):
                hash_archivo.update(bloque)

    digest = hash_archivo.hexdigest()

    return {
        "doc_id": f"DOC-{digest[:12]}",
        "sha256": digest
    }


def obtener_firma_cache(archivo, carpeta_corpus):
    estadistica = archivo.stat()

    return {
        "version": CACHE_VERSION,
        "ruta_relativa": archivo.relative_to(
            carpeta_corpus
        ).as_posix(),
        "tamano": estadistica.st_size,
        "mtime_ns": estadistica.st_mtime_ns
    }


def clave_cache(firma):
    return json.dumps(
        firma,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":")
    )


def cargar_cache_lecturas(ruta_cache):
    cache = {}

    if ruta_cache is None or not ruta_cache.exists():
        return cache

    with ruta_cache.open("r", encoding="utf-8") as archivo_cache:
        for linea in archivo_cache:
            linea = linea.strip()

            if not linea:
                continue

            try:
                registro = json.loads(linea)
                firma = registro["firma"]
                texto = registro["texto"]
                doc_id = registro["doc_id"]
                modo_doc_id = registro.get(
                    "doc_id_modo",
                    MODO_DOC_ID_CONTENIDO
                )
            except (KeyError, TypeError, json.JSONDecodeError):
                continue

            cache[clave_cache(firma)] = {
                "texto": texto,
                "doc_id": doc_id,
                "doc_id_modo": modo_doc_id
            }

    return cache


def guardar_cache_lecturas(cache, ruta_cache):
    if ruta_cache is None:
        return

    ruta_cache.parent.mkdir(
        parents=True,
        exist_ok=True
    )

    with ruta_cache.open("w", encoding="utf-8") as archivo_cache:
        for clave in sorted(cache):
            entrada = cache[clave]
            registro = {
                "firma": json.loads(clave),
                "doc_id": entrada["doc_id"],
                "doc_id_modo": entrada.get(
                    "doc_id_modo",
                    MODO_DOC_ID_RAPIDO
                ),
                "texto": entrada["texto"]
            }

            archivo_cache.write(
                json.dumps(
                    registro,
                    ensure_ascii=False
                )
                + "\n"
            )


def obtener_fenomeno(archivo, carpeta_corpus):
    ruta_relativa = archivo.relative_to(carpeta_corpus)

    for parte in ruta_relativa.parts:
        parte_normalizada = parte.upper()

        if parte_normalizada.startswith("F1_"):
            return 1

        if parte_normalizada.startswith("F2_"):
            return 2

        if parte_normalizada.startswith("F3_"):
            return 3

    return None


# =========================================================
# DETECCIÓN LIGERA DE IDIOMA
# =========================================================

PALABRAS_IDIOMA = {
    "es": {
        "el", "la", "los", "las", "de", "del", "que", "en",
        "para", "por", "con", "una", "un", "como", "se",
        "su", "sus", "entre", "sobre", "también", "desde"
    },
    "en": {
        "the", "of", "and", "to", "in", "for", "that", "with",
        "as", "is", "are", "from", "by", "this", "an", "on",
        "between", "also", "their", "its"
    },
    "pt": {
        "o", "a", "os", "as", "de", "do", "da", "que", "em",
        "para", "por", "com", "uma", "um", "como", "se",
        "sua", "suas", "entre", "sobre", "também", "dos"
    }
}


def quitar_tildes(texto):
    texto_normalizado = unicodedata.normalize("NFD", texto)

    return "".join(
        caracter
        for caracter in texto_normalizado
        if unicodedata.category(caracter) != "Mn"
    )


def detectar_idioma(texto):
    # Limitar la muestra evita procesar documentos enormes completos.
    muestra = texto[:20000].lower()

    palabras = re.findall(
        r"[a-záéíóúüñãõç]+",
        muestra
    )

    if not palabras:
        return "desconocido"

    puntuaciones = {
        "es": 0,
        "en": 0,
        "pt": 0
    }

    for palabra in palabras:
        palabra_sin_tilde = quitar_tildes(palabra)

        for idioma, vocabulario in PALABRAS_IDIOMA.items():
            vocabulario_normalizado = {
                quitar_tildes(elemento)
                for elemento in vocabulario
            }

            if palabra_sin_tilde in vocabulario_normalizado:
                puntuaciones[idioma] += 1

    idioma_ganador = max(
        puntuaciones,
        key=puntuaciones.get
    )

    puntuacion_maxima = puntuaciones[idioma_ganador]

    if puntuacion_maxima < 3:
        return "desconocido"

    puntuaciones_ordenadas = sorted(
        puntuaciones.values(),
        reverse=True
    )

    # Evitar decidir cuando dos idiomas quedan prácticamente empatados.
    if (
        len(puntuaciones_ordenadas) > 1
        and puntuaciones_ordenadas[0]
        == puntuaciones_ordenadas[1]
    ):
        return "desconocido"

    return idioma_ganador


# =========================================================
# LECTORES
# =========================================================

def leer_txt(archivo):
    with open(
        archivo,
        encoding="utf-8"
    ) as archivo_txt:
        return archivo_txt.read()


def extraer_texto_json(datos, ruta=""):
    fragmentos = []

    if isinstance(datos, dict):
        for clave, valor in datos.items():
            ruta_actual = (
                f"{ruta}.{clave}"
                if ruta
                else str(clave)
            )

            fragmentos.extend(
                extraer_texto_json(
                    valor,
                    ruta_actual
                )
            )

    elif isinstance(datos, list):
        for indice, elemento in enumerate(datos):
            ruta_actual = (
                f"{ruta}[{indice}]"
                if ruta
                else f"[{indice}]"
            )

            fragmentos.extend(
                extraer_texto_json(
                    elemento,
                    ruta_actual
                )
            )

    elif isinstance(datos, str):
        valor = datos.strip()

        if valor and not _es_url_pura(valor):
            if ruta:
                fragmentos.append(f"{ruta}: {valor}")
            else:
                fragmentos.append(valor)

    elif isinstance(datos, (int, float, bool)):
        if ruta:
            fragmentos.append(f"{ruta}: {datos}")
        else:
            fragmentos.append(str(datos))

    return fragmentos


def _extraer_urls_json(datos):
    urls = []

    def recorrer(valor):
        if isinstance(valor, dict):
            for elemento in valor.values():
                recorrer(elemento)

        elif isinstance(valor, list):
            for elemento in valor:
                recorrer(elemento)

        elif isinstance(valor, str):
            for coincidencia in PATRON_URL.findall(valor):
                url = _normalizar_url(coincidencia)

                if url:
                    urls.append(url)

    recorrer(datos)

    return sorted(set(urls))


def extraer_metadata_json(archivo):
    with open(
        archivo,
        encoding="utf-8-sig"
    ) as archivo_json:
        datos = json.load(archivo_json)

    urls = _extraer_urls_json(datos)

    if not urls:
        return {}

    return {
        "procedencia": {
            "urls": urls
        }
    }


def leer_json(archivo):
    with open(
        archivo,
        encoding="utf-8-sig"
    ) as archivo_json:
        datos = json.load(archivo_json)

    fragmentos = extraer_texto_json(datos)

    return "\n".join(fragmentos)


def leer_csv(archivo):
    filas_texto = []

    with open(
        archivo,
        encoding="utf-8-sig",
        newline=""
    ) as archivo_csv:

        lector = csv.DictReader(archivo_csv)

        if not lector.fieldnames:
            raise ValueError(
                "El archivo CSV no tiene encabezados válidos"
            )

        for fila in lector:
            campos = []

            for columna, valor in fila.items():
                if columna is None or valor is None:
                    continue

                columna = columna.strip()
                valor = valor.strip()

                if not columna or not valor:
                    continue

                campos.append(f"{columna}: {valor}")

            if campos:
                filas_texto.append(" | ".join(campos))

    return "\n".join(filas_texto)


def leer_xlsx(archivo):
    filas_texto = []

    libro = load_workbook(
        archivo,
        read_only=True,
        data_only=True
    )

    try:
        for hoja in libro.worksheets:
            filas = hoja.iter_rows(values_only=True)

            encabezados = None

            for fila in filas:
                valores = list(fila)

                if not any(
                    valor is not None
                    and str(valor).strip()
                    for valor in valores
                ):
                    continue

                # La primera fila no vacía se toma como encabezado.
                if encabezados is None:
                    encabezados = []

                    for indice, valor in enumerate(
                        valores,
                        start=1
                    ):
                        if valor is None or not str(valor).strip():
                            encabezados.append(
                                f"columna_{indice}"
                            )
                        else:
                            encabezados.append(
                                str(valor).strip()
                            )

                    continue

                campos = []

                for encabezado, valor in zip(
                    encabezados,
                    valores,
                    strict=False,
                ):
                    if valor is None:
                        continue

                    valor_texto = str(valor).strip()

                    if not valor_texto:
                        continue

                    campos.append(
                        f"{encabezado}: {valor_texto}"
                    )

                if campos:
                    filas_texto.append(
                        f"[Hoja: {hoja.title}] "
                        + " | ".join(campos)
                    )

    finally:
        libro.close()

    return "\n".join(filas_texto)


def leer_html(archivo):
    if BeautifulSoup is None:
        raise ValueError(
            "Dependencia no disponible: beautifulsoup4"
        )

    with open(
        archivo,
        encoding="utf-8"
    ) as archivo_html:
        contenido = archivo_html.read()

    sopa = BeautifulSoup(
        contenido,
        "html.parser"
    )

    elementos_eliminables = [
        "script",
        "style",
        "nav",
        "noscript",
        "template",
        "svg",
        "canvas"
    ]

    for elemento in sopa(elementos_eliminables):
        elemento.decompose()

    contenido_principal = sopa.find("main")

    if contenido_principal is None:
        contenido_principal = sopa.find("article")

    if contenido_principal is None:
        contenido_principal = sopa.body

    if contenido_principal is None:
        contenido_principal = sopa

    return contenido_principal.get_text(
        separator="\n",
        strip=True
    )


def _validar_ocr_disponible():
    if Image is None:
        raise ValueError(
            "Dependencia no disponible: Pillow; "
            "no se puede ejecutar OCR"
        )

    if pytesseract is None:
        raise ValueError(
            "OCR no disponible: instale pytesseract y Tesseract"
        )


def _ocr_imagen(imagen):
    _validar_ocr_disponible()

    if TESSDATA_LOCAL.exists():
        os.environ["TESSDATA_PREFIX"] = str(TESSDATA_LOCAL)

    imagen_ocr = imagen.copy()

    if imagen_ocr.mode not in {"RGB", "L"}:
        imagen_ocr = imagen_ocr.convert("RGB")

    try:
        return pytesseract.image_to_string(
            imagen_ocr,
            lang=OCR_IDIOMAS
        )

    except TesseractNotFoundError as error:
        raise ValueError(
            "OCR no disponible: no se encontro el ejecutable "
            "de Tesseract"
        ) from error

    except TesseractError as error:
        mensaje = str(error)

        if "Error opening data file" in mensaje:
            raise ValueError(
                "Idioma OCR no instalado: se requieren "
                "spa, eng, por y chi_sim"
            ) from error

        raise ValueError(
            f"Fallo de OCR: {mensaje}"
        ) from error


def _renderizar_pagina_pdf(pagina):
    matriz = fitz.Matrix(
        OCR_DPI / 72,
        OCR_DPI / 72
    )

    pixmap = pagina.get_pixmap(
        matrix=matriz,
        alpha=False
    )

    try:
        imagen = Image.frombytes(
            "RGB",
            (pixmap.width, pixmap.height),
            pixmap.samples
        )

    finally:
        pixmap = None

    return imagen


def leer_pdf_con_ocr(archivo):
    paginas = []
    errores_pagina = []

    if fitz is None:
        raise ValueError(
            "Dependencia no disponible: PyMuPDF"
        )

    _validar_ocr_disponible()

    with fitz.open(archivo) as documento:
        if documento.needs_pass:
            raise ValueError(
                "PDF protegido: el PDF esta protegido con contrasena"
            )

        if documento.page_count == 0:
            raise ValueError(
                "El PDF no contiene paginas"
            )

        for numero_pagina, pagina in enumerate(
            documento,
            start=1
        ):
            try:
                with _renderizar_pagina_pdf(pagina) as imagen:
                    texto_pagina = limpiar_texto(
                        _ocr_imagen(imagen)
                    )

                if texto_pagina:
                    paginas.append(
                        f"[Pagina {numero_pagina}]\n"
                        f"{texto_pagina}"
                    )

            except ValueError as error:
                errores_pagina.append(
                    f"pagina {numero_pagina}: {error}"
                )
                print(
                    f"{archivo.name}: fallo OCR en pagina "
                    f"{numero_pagina}: {error}"
                )

    if not paginas:
        detalle = "; ".join(errores_pagina)

        if detalle:
            raise ValueError(
                "PDF escaneado sin texto recuperable por OCR; "
                f"{detalle}"
            )

        raise ValueError(
            "PDF escaneado sin texto recuperable por OCR"
        )

    return "\n\n".join(paginas)


def _leer_ocr_pagina_pdf(pagina, numero_pagina, archivo):
    try:
        with _renderizar_pagina_pdf(pagina) as imagen:
            texto_pagina = limpiar_texto(
                _ocr_imagen(imagen)
            )

        if texto_pagina:
            return (
                f"[Pagina {numero_pagina}]\n"
                f"{texto_pagina}"
            )

        return ""

    except ValueError as error:
        raise ValueError(
            f"{archivo.name}: fallo OCR en pagina "
            f"{numero_pagina}: {error}"
        ) from error


def _texto_nativo_util(texto: str) -> bool:
    return bool(texto.strip())


def leer_pdf(archivo, usar_ocr=True, diagnostico=None):
    paginas = []
    paginas_sin_texto = []
    paginas_con_texto = 0

    if fitz is None:
        raise ValueError(
            "Dependencia no disponible: PyMuPDF"
        )

    with fitz.open(archivo) as documento:
        if documento.needs_pass:
            raise ValueError(
                "PDF protegido: el PDF esta protegido con contrasena"
            )

        if documento.page_count == 0:
            raise ValueError(
                "El PDF no contiene paginas"
            )

        if diagnostico is not None:
            diagnostico.clear()
            diagnostico.update({
                "paginas": documento.page_count,
                "paginas_nativas": 0,
                "paginas_ocr": 0,
                "paginas_sin_texto": 0
            })

        for numero_pagina, pagina in enumerate(
            documento,
            start=1
        ):
            texto_pagina = pagina.get_text(
                "text"
            )

            if _texto_nativo_util(texto_pagina):
                paginas_con_texto += 1
                texto_pagina = texto_pagina.strip()
                paginas.append(
                    f"[Pagina {numero_pagina}]\n"
                    f"{texto_pagina}"
                )

                if diagnostico is not None:
                    diagnostico["paginas_nativas"] += 1
            else:
                paginas.append("")
                paginas_sin_texto.append(
                    (numero_pagina, pagina, len(paginas) - 1)
                )

        if paginas_sin_texto:
            if not usar_ocr:
                if diagnostico is not None:
                    diagnostico["paginas_sin_texto"] = len(
                        paginas_sin_texto
                    )

                if paginas_con_texto == 0:
                    raise ValueError(
                        "PDF sin texto extraible; OCR desactivado"
                    )

                raise ValueError(
                    "PDF con paginas sin texto extraible; "
                    "OCR desactivado"
                )

            print(
                f"{archivo.name}: "
                f"{len(paginas_sin_texto)} pagina(s) "
                "sin texto extraible; aplicando OCR solo alli"
            )

            for numero_pagina, pagina, indice_pagina in paginas_sin_texto:
                texto_ocr = _leer_ocr_pagina_pdf(
                    pagina,
                    numero_pagina,
                    archivo
                )

                if texto_ocr:
                    paginas[indice_pagina] = texto_ocr

                    if diagnostico is not None:
                        diagnostico["paginas_ocr"] += 1
                elif diagnostico is not None:
                    diagnostico["paginas_sin_texto"] += 1

    resultado = "\n\n".join(
        pagina
        for pagina in paginas
        if pagina
    )

    if not resultado and usar_ocr:
        raise ValueError(
            "PDF escaneado sin texto recuperable por OCR"
        )

    return resultado


def leer_imagen(archivo):
    _validar_ocr_disponible()

    try:
        with Image.open(archivo) as imagen:
            imagen.load()
            texto = _ocr_imagen(imagen)

    except UnidentifiedImageError as error:
        if archivo.suffix.lower() == ".avif":
            raise ValueError(
                "Imagen no compatible: AVIF no soportado "
                "por Pillow en este entorno"
            ) from error

        raise ValueError(
            "Imagen no compatible o corrupta"
        ) from error

    except OSError as error:
        raise ValueError(
            f"Fallo de decodificacion de imagen: {error}"
        ) from error

    return limpiar_texto(texto)


def extraer_metadata_imagen(archivo):
    if Image is None:
        return {}

    try:
        with Image.open(archivo) as imagen:
            return {
                "imagen": {
                    "ancho": imagen.width,
                    "alto": imagen.height,
                    "modo": imagen.mode
                }
            }

    except (UnidentifiedImageError, OSError):
        return {}


def extraer_metadata_archivo(archivo):
    extension = archivo.suffix.lower()

    if extension == ".json":
        return extraer_metadata_json(archivo)

    if extension in EXTENSIONES_IMAGEN:
        return extraer_metadata_imagen(archivo)

    return {}


def _leer_varint(datos, posicion):
    resultado = 0
    desplazamiento = 0

    while posicion < len(datos):
        byte = datos[posicion]
        posicion += 1
        resultado |= (byte & 0x7F) << desplazamiento

        if not byte & 0x80:
            return resultado, posicion

        desplazamiento += 7

        if desplazamiento > 63:
            raise ValueError("PBF invalido: varint demasiado largo")

    raise ValueError("PBF invalido: varint incompleto")


def _leer_campo(datos, posicion):
    clave, posicion = _leer_varint(datos, posicion)
    numero_campo = clave >> 3
    tipo_wire = clave & 0x07

    if tipo_wire == 0:
        valor, posicion = _leer_varint(datos, posicion)
        return numero_campo, tipo_wire, valor, posicion

    if tipo_wire == 1:
        fin = posicion + 8
        if fin > len(datos):
            raise ValueError("PBF invalido: campo fijo64 incompleto")
        return numero_campo, tipo_wire, datos[posicion:fin], fin

    if tipo_wire == 2:
        longitud, posicion = _leer_varint(datos, posicion)
        fin = posicion + longitud
        if fin > len(datos):
            raise ValueError("PBF invalido: campo de bytes incompleto")
        return numero_campo, tipo_wire, datos[posicion:fin], fin

    if tipo_wire == 5:
        fin = posicion + 4
        if fin > len(datos):
            raise ValueError("PBF invalido: campo fijo32 incompleto")
        return numero_campo, tipo_wire, datos[posicion:fin], fin

    raise ValueError(f"PBF invalido: wire type no soportado {tipo_wire}")


def _decodificar_texto(valor):
    return valor.decode("utf-8", errors="replace")


def _decodificar_sint(valor):
    return (valor >> 1) ^ -(valor & 1)


def _decodificar_valor_mvt(datos):
    posicion = 0
    valor = None

    while posicion < len(datos):
        numero_campo, tipo_wire, dato, posicion = _leer_campo(
            datos,
            posicion
        )

        if numero_campo == 1 and tipo_wire == 2:
            valor = _decodificar_texto(dato)
        elif numero_campo == 4 and tipo_wire == 0:
            valor = int(dato)
        elif numero_campo == 5 and tipo_wire == 0:
            valor = int(dato)
        elif numero_campo == 6 and tipo_wire == 0:
            valor = _decodificar_sint(dato)
        elif numero_campo == 7 and tipo_wire == 0:
            valor = bool(dato)

    return valor


def _decodificar_feature_mvt(datos):
    posicion = 0
    feature_id = None
    tags = []

    while posicion < len(datos):
        numero_campo, tipo_wire, dato, posicion = _leer_campo(
            datos,
            posicion
        )

        if numero_campo == 1 and tipo_wire == 0:
            feature_id = dato
        elif numero_campo == 2 and tipo_wire == 2:
            posicion_tags = 0
            tags = []

            while posicion_tags < len(dato):
                tag, posicion_tags = _leer_varint(
                    dato,
                    posicion_tags
                )
                tags.append(tag)

    return feature_id, tags


def _decodificar_capa_mvt(datos):
    posicion = 0
    nombre = None
    claves = []
    valores = []
    features = []

    while posicion < len(datos):
        numero_campo, tipo_wire, dato, posicion = _leer_campo(
            datos,
            posicion
        )

        if numero_campo == 1 and tipo_wire == 2:
            nombre = _decodificar_texto(dato)
        elif numero_campo == 2 and tipo_wire == 2:
            features.append(_decodificar_feature_mvt(dato))
        elif numero_campo == 3 and tipo_wire == 2:
            claves.append(_decodificar_texto(dato))
        elif numero_campo == 4 and tipo_wire == 2:
            valores.append(_decodificar_valor_mvt(dato))

    if not nombre:
        raise ValueError("PBF no reconocido: capa sin nombre")

    return nombre, claves, valores, features


def _decodificar_tile_mvt(datos):
    posicion = 0
    capas = []

    while posicion < len(datos):
        numero_campo, tipo_wire, dato, posicion = _leer_campo(
            datos,
            posicion
        )

        if numero_campo == 3 and tipo_wire == 2:
            capas.append(_decodificar_capa_mvt(dato))

    if not capas:
        raise ValueError(
            "PBF no reconocido: no contiene capas MVT"
        )

    return capas


def _propiedades_feature(claves, valores, tags):
    propiedades = {}

    for indice in range(0, len(tags) - 1, 2):
        indice_clave = tags[indice]
        indice_valor = tags[indice + 1]

        if indice_clave >= len(claves) or indice_valor >= len(valores):
            continue

        clave = claves[indice_clave]
        valor = valores[indice_valor]

        if valor is None:
            continue

        propiedades[clave] = valor

    return propiedades


def fingerprint_pbf(nombre_capa, feature_id, propiedades):
    if feature_id is not None:
        base = {
            "capa": nombre_capa,
            "id": feature_id
        }
    else:
        base = {
            "capa": nombre_capa,
            "propiedades": propiedades
        }

    normalizado = json.dumps(
        base,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":")
    )

    return hashlib.sha256(
        normalizado.encode("utf-8")
    ).hexdigest()


def leer_pbf(archivo, fingerprints=None):
    datos = archivo.read_bytes()

    if not datos:
        return ""

    try:
        capas = _decodificar_tile_mvt(datos)

    except ValueError:
        raise

    except (IndexError, UnicodeDecodeError) as error:
        raise ValueError(
            f"Fallo de decodificacion PBF: {error}"
        ) from error

    lineas = []
    for nombre_capa, claves, valores, features in sorted(
        capas,
        key=lambda capa: capa[0]
    ):
        for feature_id, tags in features:
            propiedades = _propiedades_feature(
                claves,
                valores,
                tags
            )

            if not propiedades:
                continue

            fingerprint = fingerprint_pbf(
                nombre_capa,
                feature_id,
                propiedades
            )

            if fingerprints is not None:
                if fingerprint in fingerprints:
                    continue

                fingerprints.add(fingerprint)

            campos = [
                f"{clave}: {propiedades[clave]}"
                for clave in sorted(propiedades)
            ]

            lineas.append(
                f"capa: {nombre_capa} | "
                + " | ".join(campos)
            )

    # Solo las propiedades nuevas son contenido de corpus. Un tile cuyas
    # features ya aparecieron debe quedar vacío; diagnosticar_pbf conserva
    # el detalle operativo fuera de documento.texto.
    return "\n".join(lineas)


def diagnosticar_pbf(archivo, fingerprints=None, limite_muestras=3):
    datos = archivo.read_bytes()
    diagnostico = {
        "archivo": archivo.name,
        "tamano_bytes": len(datos),
        "parser": "MVT protobuf manual",
        "estado": "EMPTY" if not datos else "UNKNOWN",
        "capas": 0,
        "features": 0,
        "features_con_propiedades": 0,
        "features_duplicadas": 0,
        "muestras": []
    }

    if not datos:
        return diagnostico

    try:
        capas = _decodificar_tile_mvt(datos)
    except ValueError as error:
        diagnostico["estado"] = "PARSE_ERROR"
        diagnostico["error"] = str(error)
        return diagnostico
    except (IndexError, UnicodeDecodeError) as error:
        diagnostico["estado"] = "PARSE_ERROR"
        diagnostico["error"] = f"Fallo de decodificacion PBF: {error}"
        return diagnostico

    diagnostico["estado"] = "VALID"
    diagnostico["capas"] = len(capas)
    capas_resumen = []

    for nombre_capa, claves, valores, features in sorted(
        capas,
        key=lambda capa: capa[0]
    ):
        resumen_capa = {
            "capa": nombre_capa,
            "features": len(features),
            "claves": len(claves),
            "valores": len(valores)
        }
        capas_resumen.append(resumen_capa)
        diagnostico["features"] += len(features)

        for feature_id, tags in features:
            propiedades = _propiedades_feature(
                claves,
                valores,
                tags
            )

            if not propiedades:
                continue

            diagnostico["features_con_propiedades"] += 1
            fingerprint = fingerprint_pbf(
                nombre_capa,
                feature_id,
                propiedades
            )

            if fingerprints is not None and fingerprint in fingerprints:
                diagnostico["features_duplicadas"] += 1

            if len(diagnostico["muestras"]) < limite_muestras:
                diagnostico["muestras"].append(
                    {
                        "capa": nombre_capa,
                        "propiedades": {
                            clave: propiedades[clave]
                            for clave in sorted(propiedades)[:8]
                        }
                    }
                )

    diagnostico["capas_resumen"] = capas_resumen[:10]

    return diagnostico


# =========================================================
# PROCESAMIENTO
# =========================================================

def leer_archivo(archivo, fingerprints_pbf=None, usar_ocr_pdf=True):
    extension = archivo.suffix.lower()

    if extension == ".pdf":
        return leer_pdf(
            archivo,
            usar_ocr=usar_ocr_pdf
        )

    if extension == ".csv":
        return leer_csv(archivo)

    if extension == ".json":
        return leer_json(archivo)

    if extension in [".html", ".htm"]:
        return leer_html(archivo)

    if extension in [".txt", ".md"]:
        return leer_txt(archivo)

    if extension == ".xlsx":
        return leer_xlsx(archivo)

    if extension in EXTENSIONES_IMAGEN:
        return leer_imagen(archivo)

    if extension == ".pbf":
        return leer_pbf(
            archivo,
            fingerprints_pbf
        )

    raise ValueError("Formato no soportado")


def registrar_error(
    errores,
    archivo,
    estado,
    mensaje,
    metadata=None
):
    registro = {
        "archivo": archivo,
        "estado": estado,
        "error": mensaje
    }

    if metadata:
        registro["metadata"] = metadata

    errores.append(registro)


def es_pendiente_ocr(mensaje):
    return (
        "OCR desactivado" in mensaje
        or "PDF sin texto extraible" in mensaje
    )


def _leer_tarea_corpus(tarea, fingerprints_pbf=None):
    archivo = tarea["archivo"]
    cache = tarea.get("cache")
    clave = tarea.get("clave_cache")
    claves_cache = [
        clave_candidata
        for clave_candidata in [
            clave,
            tarea.get("clave_cache_ocr")
        ]
        if clave_candidata
    ]
    ignorar_cache = tarea.get("ignorar_cache", False)

    try:
        clave_encontrada = None

        if cache is not None and not ignorar_cache:
            for clave_candidata in claves_cache:
                if clave_candidata in cache:
                    clave_encontrada = clave_candidata
                    break

        if cache is not None and clave_encontrada is not None:
            entrada_cache = cache[clave_encontrada]
            modo_doc_id = tarea.get("modo_doc_id", MODO_DOC_ID_RAPIDO)

            if entrada_cache.get("doc_id_modo") == modo_doc_id:
                doc_id = entrada_cache["doc_id"]
                cache_doc_id_actualizado = False
            else:
                fingerprint = generar_fingerprint_archivo(
                    archivo,
                    tarea["carpeta_corpus"],
                    modo=modo_doc_id
                )
                doc_id = fingerprint["doc_id"]
                cache_doc_id_actualizado = True

            return {
                **tarea,
                "texto": entrada_cache["texto"],
                "doc_id": doc_id,
                "metadata": extraer_metadata_archivo(archivo),
                "desde_cache": True,
                "cache_doc_id_actualizado": cache_doc_id_actualizado
            }

        texto = leer_archivo(
            archivo,
            fingerprints_pbf,
            usar_ocr_pdf=tarea.get("usar_ocr", False)
        )
        texto = limpiar_texto(texto)
        fingerprint = generar_fingerprint_archivo(
            archivo,
            tarea["carpeta_corpus"],
            modo=tarea.get("modo_doc_id", MODO_DOC_ID_RAPIDO)
        )
        metadata = extraer_metadata_archivo(archivo)

        return {
            **tarea,
            "texto": texto,
            "doc_id": fingerprint["doc_id"],
            "metadata": metadata
        }

    except ValueError as error:
        return {
            **tarea,
            "error": str(error)
        }

    except Exception as error:
        return {
            **tarea,
            "error": str(error),
            "error_inesperado": True
        }


def _registrar_resultado_procesamiento(
    resultado,
    carpeta_corpus,
    documentos,
    errores,
    ids_generados,
    estadisticas
):
    archivo = resultado["archivo"]
    ruta_relativa = resultado["ruta_relativa"]

    if "error" in resultado:
        mensaje = resultado["error"]

        if mensaje == "Formato no soportado":
            estado = "omitido"
            estadisticas["omitidos"] += 1
        elif es_pendiente_ocr(mensaje):
            estado = "pendiente_ocr"
            estadisticas["pendientes_ocr"] += 1
        else:
            estado = "fallido"
            estadisticas["fallidos"] += 1

        if resultado.get("error_inesperado"):
            print(
                f"{archivo.name}: "
                f"error durante la lectura: {mensaje}"
            )
        else:
            print(f"{archivo.name}: {mensaje}")

        registrar_error(
            errores,
            ruta_relativa,
            estado,
            mensaje
        )

        return

    texto = resultado["texto"]

    if not texto:
        metadata = resultado.get("metadata", {})
        formato = obtener_formato(resultado["extension"])

        if formato == "imagen":
            print(f"{archivo.name}: imagen sin texto util")

            registrar_error(
                errores,
                ruta_relativa,
                "sin_texto_util",
                "Imagen valida sin texto recuperable por OCR",
                metadata
            )
        else:
            print(f"{archivo.name}: texto vacío")

            registrar_error(
                errores,
                ruta_relativa,
                "vacío",
                "No se extrajo contenido textual",
                metadata
            )

        estadisticas["vacios"] += 1
        return

    doc_id = resultado["doc_id"]

    if doc_id in ids_generados:
        print(f"{archivo.name}: doc_id duplicado")

        registrar_error(
            errores,
            ruta_relativa,
            "fallido",
            f"doc_id duplicado: {doc_id}"
        )

        estadisticas["fallidos"] += 1
        return

    ids_generados.add(doc_id)

    documento = {
        "doc_id": doc_id,
        "fuente": ruta_relativa,
        "formato": obtener_formato(resultado["extension"]),
        "fenomeno": resultado["fenomeno"],
        "idioma": detectar_idioma(texto),
        "texto": texto
    }

    metadata = resultado.get("metadata")

    if metadata:
        documento["metadata"] = metadata

    documentos.append(documento)
    estadisticas["procesados"] += 1


def _actualizar_cache_desde_resultado(resultado, cache_lecturas):
    if (
        cache_lecturas is None
        or "error" in resultado
        or (
            resultado.get("desde_cache")
            and not resultado.get("cache_doc_id_actualizado")
        )
    ):
        return

    clave = resultado.get("clave_cache")

    if not clave:
        return

    cache_lecturas[clave] = {
        "texto": resultado["texto"],
        "doc_id": resultado["doc_id"],
        "doc_id_modo": resultado.get(
            "modo_doc_id",
            MODO_DOC_ID_RAPIDO
        )
    }


def _tarea_para_proceso(tarea):
    return {
        clave: valor
        for clave, valor in tarea.items()
        if clave != "cache"
    }


def _registrar_buffer_en_orden(
    buffer_resultados,
    siguiente_orden,
    cache_lecturas,
    carpeta_corpus,
    documentos,
    errores,
    ids_generados,
    estadisticas
):
    while siguiente_orden in buffer_resultados:
        resultado = buffer_resultados.pop(siguiente_orden)
        _actualizar_cache_desde_resultado(
            resultado,
            cache_lecturas
        )

        _registrar_resultado_procesamiento(
            resultado,
            carpeta_corpus,
            documentos,
            errores,
            ids_generados,
            estadisticas
        )

        siguiente_orden += 1

    return siguiente_orden


def _procesar_grupo_paralelo(
    grupo,
    max_trabajadores,
    buffer_resultados,
    usar_procesos=False
):
    trabajadores = min(
        max(1, max_trabajadores),
        len(grupo)
    )

    if trabajadores <= 1:
        for tarea in grupo:
            buffer_resultados[tarea["orden"]] = _leer_tarea_corpus(tarea)

        return

    ejecutor_clase = (
        ProcessPoolExecutor
        if usar_procesos
        else ThreadPoolExecutor
    )

    def enviar(ejecutor, tarea):
        return ejecutor.submit(
            _leer_tarea_corpus,
            _tarea_para_proceso(tarea)
            if usar_procesos
            else tarea
        )

    siguiente = 0
    futuros = []

    with ejecutor_clase(max_workers=trabajadores) as ejecutor:
        while siguiente < len(grupo) and len(futuros) < trabajadores:
            futuros.append(enviar(ejecutor, grupo[siguiente]))
            siguiente += 1

        while futuros:
            for futuro in as_completed(futuros):
                futuros.remove(futuro)
                resultado = futuro.result()
                buffer_resultados[resultado["orden"]] = resultado

                if siguiente < len(grupo):
                    futuros.append(enviar(ejecutor, grupo[siguiente]))
                    siguiente += 1

                break


def _procesar_tareas_en_orden(
    tareas,
    max_trabajadores,
    max_procesos_pesados,
    fingerprints_pbf,
    cache_lecturas,
    carpeta_corpus,
    documentos,
    errores,
    ids_generados,
    estadisticas
):
    buffer_resultados = {}
    siguiente_orden = 0
    posicion = 0

    while posicion < len(tareas):
        tarea = tareas[posicion]

        if tarea.get("clave_cache") in cache_lecturas:
            buffer_resultados[tarea["orden"]] = _leer_tarea_corpus(tarea)

            siguiente_orden = _registrar_buffer_en_orden(
                buffer_resultados,
                siguiente_orden,
                cache_lecturas,
                carpeta_corpus,
                documentos,
                errores,
                ids_generados,
                estadisticas
            )

            posicion += 1
            continue

        if tarea["extension"] == ".pbf":
            buffer_resultados[tarea["orden"]] = _leer_tarea_corpus(
                tarea,
                fingerprints_pbf
            )

            siguiente_orden = _registrar_buffer_en_orden(
                buffer_resultados,
                siguiente_orden,
                cache_lecturas,
                carpeta_corpus,
                documentos,
                errores,
                ids_generados,
                estadisticas
            )

            posicion += 1
            continue

        inicio_grupo = posicion
        usar_procesos = tarea["extension"] in EXTENSIONES_PROCESOS

        while (
            posicion < len(tareas)
            and tareas[posicion]["extension"] != ".pbf"
            and (
                tareas[posicion]["extension"] in EXTENSIONES_PROCESOS
            ) == usar_procesos
        ):
            posicion += 1

        grupo = tareas[inicio_grupo:posicion]
        trabajadores = min(
            max(1, (
                max_procesos_pesados
                if usar_procesos
                else max_trabajadores
            )),
            len(grupo)
        )

        tamano_tanda = trabajadores * 4

        for inicio_tanda in range(0, len(grupo), tamano_tanda):
            _procesar_grupo_paralelo(
                grupo[inicio_tanda:inicio_tanda + tamano_tanda],
                trabajadores,
                buffer_resultados,
                usar_procesos=usar_procesos
            )

            siguiente_orden = _registrar_buffer_en_orden(
                buffer_resultados,
                siguiente_orden,
                cache_lecturas,
                carpeta_corpus,
                documentos,
                errores,
                ids_generados,
                estadisticas
            )


def procesar_corpus(
    carpeta_corpus,
    max_trabajadores=MAX_TRABAJADORES,
    max_procesos_pesados=MAX_PROCESOS_PESADOS,
    modo_doc_id=MODO_DOC_ID_RAPIDO,
    usar_ocr=False,
    ruta_cache=RUTA_CACHE_LECTURAS,
    rutas_incluidas=None,
    ignorar_cache_lectura=False
):
    documentos = []
    errores = []
    ids_generados = set()
    fingerprints_pbf = set()
    tareas = []
    cache_lecturas = cargar_cache_lecturas(ruta_cache)

    estadisticas = {
        "encontrados": 0,
        "procesados": 0,
        "omitidos": 0,
        "vacios": 0,
        "pendientes_ocr": 0,
        "fallidos": 0
    }

    for indice, archivo in enumerate(sorted(carpeta_corpus.rglob("*"))):
        if not archivo.is_file():
            continue

        if any(
            parte in DIRECTORIOS_IGNORADOS
            for parte in archivo.relative_to(carpeta_corpus).parts
        ):
            continue

        extension = archivo.suffix.lower()
        nombre_archivo = archivo.name
        nombre_normalizado = nombre_archivo.lower()

        if (
            nombre_archivo in BASURA_SILENCIOSA
            or extension == ".tmp"
            or nombre_archivo.startswith(".~")
        ):
            continue

        ruta_relativa = archivo.relative_to(
            carpeta_corpus
        ).as_posix()
        ruta_relativa_normalizada = unicodedata.normalize(
            "NFC",
            ruta_relativa
        ).casefold()

        if (
            rutas_incluidas is not None
            and ruta_relativa not in rutas_incluidas
        ):
            continue

        estadisticas["encontrados"] += 1

        motivo_auxiliar = RUTAS_AUXILIARES.get(
            ruta_relativa_normalizada
        )

        if archivo.parent == carpeta_corpus:
            motivo_auxiliar = ARCHIVOS_AUXILIARES.get(
                nombre_normalizado,
                motivo_auxiliar
            )

        if motivo_auxiliar is not None:
            print(
                f"{archivo.name}: archivo auxiliar omitido"
            )

            registrar_error(
                errores,
                ruta_relativa,
                "omitido",
                motivo_auxiliar
            )

            estadisticas["omitidos"] += 1
            continue

        if extension not in EXTENSIONES_SOPORTADAS:
            print(
                f"{archivo.name}: formato no soportado"
            )

            registrar_error(
                errores,
                ruta_relativa,
                "omitido",
                "Formato no soportado"
            )

            estadisticas["omitidos"] += 1
            continue

        fenomeno = obtener_fenomeno(
            archivo,
            carpeta_corpus
        )

        # El índice general de Codefest no pertenece
        # necesariamente a un único fenómeno.
        if (
            fenomeno is None
            and archivo.name.lower()
            == "indice_datos_codefest.xlsx"
        ):
            print(
                f"{archivo.name}: archivo auxiliar omitido"
            )

            registrar_error(
                errores,
                ruta_relativa,
                "omitido",
                "Archivo índice auxiliar del corpus"
            )

            estadisticas["omitidos"] += 1
            continue

        if fenomeno is None:
            print(
                f"{archivo.name}: "
                "no se pudo determinar el fenómeno"
            )

            registrar_error(
                errores,
                ruta_relativa,
                "fallido",
                "No se pudo determinar el fenómeno "
                "desde la carpeta de origen"
            )

            estadisticas["fallidos"] += 1
            continue

        firma_cache = obtener_firma_cache(
            archivo,
            carpeta_corpus
        )

        if extension == ".pdf":
            firma_cache["usar_ocr"] = usar_ocr
            firma_cache_ocr = dict(firma_cache)
            firma_cache_ocr["usar_ocr"] = True
        else:
            firma_cache_ocr = None

        usar_cache = extension != ".pbf"

        tarea = {
            "indice": indice,
            "orden": len(tareas),
            "archivo": archivo,
            "carpeta_corpus": carpeta_corpus,
            "ruta_relativa": ruta_relativa,
            "extension": extension,
            "fenomeno": fenomeno,
            "cache": cache_lecturas if usar_cache else None,
            "firma_cache": firma_cache if usar_cache else None,
            "clave_cache": clave_cache(firma_cache) if usar_cache else None,
            "clave_cache_ocr": (
                clave_cache(firma_cache_ocr)
                if usar_cache and firma_cache_ocr is not None and not usar_ocr
                else None
            ),
            "modo_doc_id": modo_doc_id,
            "usar_ocr": usar_ocr,
            "ignorar_cache": ignorar_cache_lectura
        }

        tareas.append(tarea)

    total_lecturas = len(tareas)

    if total_lecturas:
        print(
            "Iniciando lectura del corpus...",
            flush=True
        )

    _procesar_tareas_en_orden(
        tareas,
        max_trabajadores,
        max_procesos_pesados,
        fingerprints_pbf,
        cache_lecturas,
        carpeta_corpus,
        documentos,
        errores,
        ids_generados,
        estadisticas
    )

    guardar_cache_lecturas(
        cache_lecturas,
        ruta_cache
    )

    return documentos, errores, estadisticas


# =========================================================
# GUARDADO
# =========================================================

def guardar_jsonl(registros, ruta):
    with open(
        ruta,
        "w",
        encoding="utf-8"
    ) as archivo_salida:

        for registro in registros:
            linea = json.dumps(
                registro,
                ensure_ascii=False
            )

            archivo_salida.write(linea + "\n")


def cargar_jsonl(ruta):
    if not ruta.exists():
        return []

    registros = []

    with ruta.open("r", encoding="utf-8") as archivo:
        for linea in archivo:
            linea = linea.strip()

            if linea:
                registros.append(json.loads(linea))

    return registros


def cargar_rutas_pendientes_ocr(ruta_errores):
    pendientes = set()

    for error in cargar_jsonl(ruta_errores):
        if error.get("estado") == "pendiente_ocr":
            archivo = error.get("archivo")

            if isinstance(archivo, str) and archivo:
                pendientes.add(archivo)

    return pendientes


def fusionar_por_fuente(documentos_base, documentos_nuevos):
    fusionados = {
        documento["fuente"]: documento
        for documento in documentos_base
    }

    for documento in documentos_nuevos:
        fusionados[documento["fuente"]] = documento

    return [
        fusionados[fuente]
        for fuente in sorted(fusionados)
    ]


def recalcular_estadisticas_finales(
    estadisticas,
    documentos,
    errores
):
    estadisticas["procesados"] = len(documentos)
    estadisticas["omitidos"] = sum(
        1
        for error in errores
        if error.get("estado") == "omitido"
    )
    estadisticas["vacios"] = sum(
        1
        for error in errores
        if error.get("estado") in {"vacío", "vacÃ­o", "vacÃƒÂ­o"}
    )
    estadisticas["pendientes_ocr"] = sum(
        1
        for error in errores
        if error.get("estado") == "pendiente_ocr"
    )
    estadisticas["fallidos"] = sum(
        1
        for error in errores
        if error.get("estado") == "fallido"
    )


def resolver_pendientes_ocr(
    carpeta_corpus,
    documentos_base,
    errores_base,
    rutas_pendientes_ocr,
    estadisticas_base,
    max_trabajadores,
    max_procesos_pesados,
    modo_doc_id,
    ruta_cache
):
    print(
        "Procesando pendientes OCR: "
        f"{len(rutas_pendientes_ocr)} archivo(s)",
        flush=True
    )

    documentos_ocr, errores_ocr, _estadisticas_ocr = procesar_corpus(
        carpeta_corpus,
        max_trabajadores=max_trabajadores,
        max_procesos_pesados=max_procesos_pesados,
        modo_doc_id=modo_doc_id,
        usar_ocr=True,
        ruta_cache=ruta_cache,
        rutas_incluidas=rutas_pendientes_ocr,
        ignorar_cache_lectura=True
    )

    documentos = fusionar_por_fuente(
        documentos_base,
        documentos_ocr
    )

    errores = [
        error
        for error in errores_base
        if error.get("estado") != "pendiente_ocr"
    ]
    errores.extend(errores_ocr)

    estadisticas = dict(estadisticas_base)
    recalcular_estadisticas_finales(
        estadisticas,
        documentos,
        errores
    )

    print(
        "Pendientes OCR resueltos: "
        f"{len(documentos_ocr)}/{len(rutas_pendientes_ocr)}",
        flush=True
    )

    return documentos, errores, estadisticas


def mostrar_estadisticas(
    documentos,
    errores,
    estadisticas,
    ruta_documentos,
    ruta_errores
):
    print(
        f"\nSe guardaron {len(documentos)} documentos "
        f"en {ruta_documentos}"
    )

    print(
        f"Se registraron {len(errores)} incidencias "
        f"en {ruta_errores}"
    )

    print("\nResumen del procesamiento")
    print(
        f"Archivos encontrados: "
        f"{estadisticas['encontrados']}"
    )
    print(
        f"Procesados correctamente: "
        f"{estadisticas['procesados']}"
    )
    print(
        f"Omitidos: {estadisticas['omitidos']}"
    )
    print(
        f"Vacíos: {estadisticas['vacios']}"
    )
    print(
        f"Pendientes OCR: {estadisticas['pendientes_ocr']}"
    )
    print(
        f"Fallidos: {estadisticas['fallidos']}"
    )


# =========================================================
# EJECUCIÓN PRINCIPAL
# =========================================================

def construir_parser():
    parser = argparse.ArgumentParser(
        description="Procesa el corpus documental NERV a JSON Lines."
    )

    parser.add_argument(
        "--corpus",
        default=str(RAIZ_PROYECTO / "corpus" / "raw"),
        help="Carpeta raiz del corpus a procesar."
    )

    parser.add_argument(
        "--salida",
        default=str(RAIZ_PROYECTO / "corpus" / "processed"),
        help="Carpeta donde se guardan documentos.jsonl y errores.jsonl."
    )

    parser.add_argument(
        "--trabajadores",
        type=int,
        default=MAX_TRABAJADORES,
        help=(
            "Numero maximo de archivos leidos en paralelo. "
            "Use 1 para desactivar paralelismo."
        )
    )

    parser.add_argument(
        "--procesos-pesados",
        type=int,
        default=MAX_PROCESOS_PESADOS,
        help=(
            "Numero maximo de procesos para PDF e imagenes. "
            "Use 1 si el equipo se queda corto de memoria."
        )
    )

    parser.add_argument(
        "--sin-cache",
        action="store_true",
        help="Desactiva el cache persistente de lecturas."
    )

    parser.add_argument(
        "--ocr",
        action="store_true",
        help="Activa OCR como fallback durante la primera pasada de PDFs."
    )

    parser.add_argument(
        "--sin-ocr",
        action="store_true",
        help=(
            "Ejecuta solo la pasada rapida y deja PDFs escaneados "
            "como pendiente_ocr."
        )
    )

    parser.add_argument(
        "--solo-pendientes-ocr",
        action="store_true",
        help=(
            "Procesa con OCR solo las rutas marcadas como pendiente_ocr "
            "en errores.jsonl."
        )
    )

    parser.add_argument(
        "--doc-id-contenido",
        action="store_true",
        help=(
            "Genera doc_id leyendo el contenido completo del archivo. "
            "Es mas lento, pero independiente de la fecha de modificacion."
        )
    )

    parser.add_argument(
        "--diagnosticar-pbf",
        help="Ruta de un archivo .pbf para imprimir diagnostico compacto."
    )

    return parser


def main(argv=None):
    argumentos = construir_parser().parse_args(argv)
    carpeta_corpus = Path(argumentos.corpus).expanduser().resolve()
    carpeta_salida = Path(argumentos.salida).expanduser().resolve()

    if argumentos.diagnosticar_pbf:
        ruta_pbf = Path(argumentos.diagnosticar_pbf).expanduser().resolve()
        diagnostico = diagnosticar_pbf(ruta_pbf)
        print(
            json.dumps(
                diagnostico,
                ensure_ascii=False,
                indent=2
            )
        )
        return

    if not carpeta_corpus.exists():
        print(
            "Error: la carpeta del corpus no existe"
        )
        sys.exit(1)

    if not carpeta_corpus.is_dir():
        print(
            "Error: la ruta del corpus "
            "no corresponde a una carpeta"
        )
        sys.exit(1)

    carpeta_salida.mkdir(
        parents=True,
        exist_ok=True
    )

    ruta_documentos = (
        carpeta_salida / "documentos.jsonl"
    )

    ruta_errores = (
        carpeta_salida / "errores.jsonl"
    )

    rutas_pendientes_ocr = None
    max_trabajadores = max(1, argumentos.trabajadores)
    max_procesos_pesados = max(1, argumentos.procesos_pesados)
    modo_doc_id = (
        MODO_DOC_ID_CONTENIDO
        if argumentos.doc_id_contenido
        else MODO_DOC_ID_RAPIDO
    )
    ruta_cache = (
        None
        if argumentos.sin_cache
        else RUTA_CACHE_LECTURAS
    )

    if argumentos.solo_pendientes_ocr:
        rutas_pendientes_ocr = cargar_rutas_pendientes_ocr(
            ruta_errores
        )
        documentos_base = cargar_jsonl(ruta_documentos)
        errores_base = cargar_jsonl(ruta_errores)

        if not rutas_pendientes_ocr:
            print(
                "No hay pendientes OCR en errores.jsonl"
            )
            return

        estadisticas_base = {
            "encontrados": len(documentos_base) + len(errores_base),
            "procesados": len(documentos_base),
            "omitidos": 0,
            "vacios": 0,
            "pendientes_ocr": len(rutas_pendientes_ocr),
            "fallidos": 0
        }

        documentos, errores, estadisticas = resolver_pendientes_ocr(
            carpeta_corpus,
            documentos_base,
            errores_base,
            rutas_pendientes_ocr,
            estadisticas_base,
            max_trabajadores,
            max_procesos_pesados,
            modo_doc_id,
            ruta_cache
        )

    else:
        documentos, errores, estadisticas = procesar_corpus(
            carpeta_corpus,
            max_trabajadores=max_trabajadores,
            max_procesos_pesados=max_procesos_pesados,
            modo_doc_id=modo_doc_id,
            usar_ocr=argumentos.ocr,
            ruta_cache=ruta_cache
        )

        rutas_pendientes_ocr = {
            error["archivo"]
            for error in errores
            if error.get("estado") == "pendiente_ocr"
        }

        if (
            rutas_pendientes_ocr
            and not argumentos.sin_ocr
            and not argumentos.ocr
        ):
            documentos, errores, estadisticas = resolver_pendientes_ocr(
                carpeta_corpus,
                documentos,
                errores,
                rutas_pendientes_ocr,
                estadisticas,
                max_trabajadores,
                max_procesos_pesados,
                modo_doc_id,
                ruta_cache
            )

    guardar_jsonl(
        documentos,
        ruta_documentos
    )

    guardar_jsonl(
        errores,
        ruta_errores
    )

    mostrar_estadisticas(
        documentos,
        errores,
        estadisticas,
        ruta_documentos,
        ruta_errores
    )


if __name__ == "__main__":
    main()
